"""
Converts the factory Excel files into JSON for the NOS import endpoints.

  python build_import.py --yarn "Yarn Stock Summary.xlsx" --program "Machine wise Knitting Program -2026.xlsx" --as-of 2026-09-28

Writes yarn-import.json  (POST /api/yarn/import, admin)
       knitting-import.json (POST /api/knitting/import, admin)
"""
import argparse, datetime, json, re
import openpyxl

def iso(v):
    if isinstance(v, datetime.datetime): return '' if v.year < 2000 else v.strftime('%Y-%m-%d')
    s = str(v or '').strip(); m = re.match(r'^(\d{1,2})[./](\d{1,2})[./](\d{2,4})$', s)
    if m: y = m.group(3); y = '20' + y if len(y) == 2 else y; return '%s-%02d-%02d' % (y, int(m.group(2)), int(m.group(1)))
    return ''
def num(v):
    try: return round(float(v), 2)
    except Exception: return 0.0
def guess_count(desc, codes):
    t = str(desc).upper(); m = re.search(r'(\d+)\s*D\b', t)
    if m and (m.group(1) + 'D') in codes: return m.group(1) + 'D'
    m = re.search(r'(\d+)\s*(?:S\b|/1|\'S)', t)
    if not m: return None
    base = m.group(1) + '/1'; tag = 'ORG' if 'ORG' in t else 'BCI' if 'BCI' in t else 'CVC' if 'CVC' in t else 'GM' if ('MEL' in t or 'GM' in t) else 'VIS' if 'VIS' in t else ''
    c = [x for x in codes if x.upper().startswith(base)]
    best = next((x for x in c if tag and tag in x.upper()), None) or next((x for x in c if re.search(r'\bW\b', x.upper())), None) or (c[0] if c else None)
    return best

def yarn(path, as_of):
    wb = openpyxl.load_workbook(path, data_only=True)
    counts = []
    for r in wb['safety stock'].iter_rows(min_row=2, values_only=True):
        if r[0] in (None, '', 'Yarn Cout', 'Total'): continue
        counts.append({'code': str(r[0]).strip(), 'moqKg': num(r[2])})
    codes = [c['code'] for c in counts]
    lots, started = [], False
    for r in wb['Yarn Stock'].iter_rows(values_only=True):
        if not started:
            started = any(str(c or '').strip().startswith('Description of') for c in r); continue
        desc = str(r[2] or '').strip()
        if not desc or desc.lower().startswith(('sub total', 'grand', 'total')): continue
        lots.append({'yarnDesc': desc, 'countCode': guess_count(desc, codes), 'supplier': str(r[3] or '').strip() or None, 'lotNo': str(r[4] or '').strip().replace('.0', '') or '-', 'lpNo': str(r[14] or '').strip().replace('.0', ''),
                     'receivedAt': iso(r[5]) or None, 'lcNo': str(r[6] or '').strip() or None, 'location': str(r[15] or '').strip() or None, 'closingKg': num(r[13])})
    # the same yarn + lot + L.P no. can appear twice in the sheet: add the kg together
    merged = {}
    for l in lots:
        k = (l['yarnDesc'], l['lotNo'], l['lpNo'])
        if k in merged: merged[k]['closingKg'] = round(merged[k]['closingKg'] + l['closingKg'], 2)
        else: merged[k] = l
    return {'asOf': as_of, 'counts': counts, 'lots': list(merged.values())}

def program(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    machines = [{'id': 'M-%02d' % int(r[1]), 'dia': str(r[0]).strip(), 'capKgDay': num(r[25])} for r in wb['Summary'].iter_rows(min_row=3, values_only=True) if isinstance(r[1], (int, float)) and num(r[25]) > 0]
    sheet = [s for s in wb.sheetnames if re.match(r'^[A-Z][a-z]{2}-\d{2}$', s)][-1]
    prog = []
    for r in wb[sheet].iter_rows(min_row=3, values_only=True):
        if not isinstance(r[2], (int, float)) or r[4] in (None, '') or r[5] in (None, ''): continue
        pos = str(r[21] if r[21] is not None else '').strip().upper().replace('.0', '')
        if pos in ('CLOSED', 'STOP') or num(r[15]) - num(r[17]) <= 0: continue
        prog.append({'mc': 'M-%02d' % int(r[2]), 'dia': str(r[3] or '').strip(), 'buyer': str(r[4]).strip(), 'job': str(r[5]).strip(), 'ta': iso(r[6]) or None, 'count': str(r[7] or '').strip(), 'brand': str(r[8] or '').strip(), 'lot': str(r[9] or '').strip().replace('.0', ''),
                     'fabric': str(r[10] or '').strip() or '-', 'gsm': str(r[11] or '').replace('.0', ''), 'fwid': str(r[12] or '').replace('.0', ''), 'sl': str(r[13] or '').strip(), 'colour': str(r[14] or '').strip(),
                     'order': num(r[15]), 'prod': num(r[17]), 'prog': str(r[19] or '').replace('.0', ''), 'progRcv': iso(r[20]) or None, 'pos': pos, 'remark': str(r[22] or '').strip(), 'ship': iso(r[23]) or None})
    # machines used in the program but missing from the Summary sheet: same dia & gauge average capacity, reported
    norm = lambda d: '-'.join(re.findall(r'\d+', str(d))[:2])
    known = {m['id'] for m in machines}; added = []
    for p in prog:
        if p['mc'] in known: continue
        same = [m['capKgDay'] for m in machines if norm(m['dia']) == norm(p['dia'])]
        cap = round(sum(same) / len(same)) if same else round(sum(m['capKgDay'] for m in machines) / len(machines))
        machines.append({'id': p['mc'], 'dia': p['dia'].replace("''", '"'), 'capKgDay': cap}); known.add(p['mc']); added.append(f"{p['mc']} {p['dia']} assumed {cap} kg/day")
    for p in prog: p.pop('dia', None)
    if added: print('machines missing in Summary, capacity assumed:', '; '.join(added))
    return {'machines': machines, 'program': prog}

if __name__ == '__main__':
    a = argparse.ArgumentParser(); a.add_argument('--yarn'); a.add_argument('--program'); a.add_argument('--as-of', required=True); o = a.parse_args()
    if o.yarn: d = yarn(o.yarn, o.as_of); json.dump(d, open('yarn-import.json', 'w'), indent=1); print('yarn:', len(d['counts']), 'counts,', len(d['lots']), 'lots')
    if o.program: d = program(o.program); json.dump(d, open('knitting-import.json', 'w'), indent=1); print('knitting:', len(d['machines']), 'machines,', len(d['program']), 'open program lines')
